#!/usr/bin/env python3
"""Build a source-backed scorecard workbook from the 投早投小 template."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "投早投小打分卡.xlsx"


EXTRA_HEADERS = {
    8: "信息来源",
    9: "证据等级",
    10: "依据摘录/链接",
    11: "待核实事项",
}


def normalize_sources(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(x) for x in value if x)
    return str(value)


def ensure_columns(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, header in EXTRA_HEADERS.items():
        cell = ws.cell(row=2, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 18,
        "B": 20,
        "C": 24,
        "D": 58,
        "E": 10,
        "F": 42,
        "G": 42,
        "H": 42,
        "I": 16,
        "J": 46,
        "K": 36,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=11):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def fill_items(ws, items: list[dict[str, Any]]) -> None:
    for item in items:
        row = int(item["row"])
        if "score" in item:
            ws.cell(row=row, column=5).value = item["score"]
        if "description" in item:
            ws.cell(row=row, column=6).value = item["description"]
        if "indicator_note" in item:
            ws.cell(row=row, column=7).value = item["indicator_note"]
        ws.cell(row=row, column=8).value = normalize_sources(item.get("sources"))
        ws.cell(row=row, column=9).value = item.get("evidence_level", "")
        ws.cell(row=row, column=10).value = normalize_sources(item.get("quote_or_link"))
        ws.cell(row=row, column=11).value = item.get("to_verify", "")


def add_summary_sheet(wb, data: dict[str, Any]) -> None:
    if "项目摘要" in wb.sheetnames:
        del wb["项目摘要"]
    ws = wb.create_sheet("项目摘要", 0)
    rows = [
        ("项目名称", data.get("project_name", "")),
        ("评分日期/口径", data.get("as_of_date", "")),
        ("综合结论", data.get("conclusion", "")),
        ("建议动作", data.get("recommendation", "")),
        ("总分", data.get("total_score", "")),
        ("主要加分项", normalize_sources(data.get("key_positives"))),
        ("主要扣分/红旗", normalize_sources(data.get("key_risks"))),
        ("关键待核实事项", normalize_sources(data.get("key_to_verify"))),
    ]
    ws.append(["字段", "内容"])
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")


def add_sources_sheet(wb, sources: list[dict[str, Any]]) -> None:
    if "来源清单" in wb.sheetnames:
        del wb["来源清单"]
    ws = wb.create_sheet("来源清单")
    ws.append(["编号", "类型", "完整引用", "URL/文件定位", "备注"])
    for source in sources:
        ws.append([
            source.get("id", ""),
            source.get("type", ""),
            source.get("citation", ""),
            source.get("url_or_locator", ""),
            source.get("notes", ""),
        ])
    widths = {"A": 12, "B": 18, "C": 70, "D": 60, "E": 36}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build source-backed scorecard workbook.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Path to template .xlsx")
    parser.add_argument("--answers", help="Optional JSON answers file")
    parser.add_argument("--output", required=True, help="Output .xlsx path")
    args = parser.parse_args()

    data: dict[str, Any] = {}
    if args.answers:
        data = json.loads(Path(args.answers).read_text(encoding="utf-8"))

    wb = load_workbook(args.template)
    ws = wb[wb.sheetnames[0]]
    ensure_columns(ws)
    if data.get("project_name"):
        ws["A1"] = f"{data['project_name']} - 投早投小科创企业打分卡"
    if data.get("items"):
        fill_items(ws, data["items"])
    add_summary_sheet(wb, data)
    add_sources_sheet(wb, data.get("sources", []))

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(output)


if __name__ == "__main__":
    main()
