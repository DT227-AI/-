#!/usr/bin/env python3
"""Build source-backed project management interview minutes from a DOCX template."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "附件1：项目管理层访谈纪要.docx"
DEFAULT_FIELD_MAP = SKILL_DIR / "references" / "field-map.csv"


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "；".join(str(x) for x in value if x)
    return str(value)


def load_field_map(path: Path) -> dict[str, dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return {row["field_id"]: row for row in csv.DictReader(f)}


def normalize_fields(raw: Any) -> dict[str, dict[str, Any]]:
    if raw is None:
        return {}
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, list):
        result: dict[str, dict[str, Any]] = {}
        for item in raw:
            fid = item.get("field_id")
            if fid:
                result[fid] = item
        return result
    raise TypeError("fields must be a dict or list")


def set_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    lines = str(text).split("\n")
    run = paragraph.add_run(lines[0] if lines else "")
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def append_visible_source_block(value: Any, item: dict[str, Any]) -> str:
    text = normalize_text(value)
    sources = normalize_text(item.get("sources"))
    evidence = normalize_text(item.get("evidence_level"))
    verify = normalize_text(item.get("to_verify"))
    parts = [text]
    if sources:
        parts.append(f"【信息来源】{sources}")
    if evidence:
        parts.append(f"【证据等级】{evidence}")
    if verify:
        parts.append(f"【待核实】{verify}")
    return "\n".join(part for part in parts if part)


def comment_text(item: dict[str, Any]) -> str:
    parts = []
    for label, key in [
        ("信息来源", "sources"),
        ("证据等级", "evidence_level"),
        ("依据摘录/链接", "quote_or_link"),
        ("待核实事项", "to_verify"),
    ]:
        value = normalize_text(item.get(key))
        if value:
            parts.append(f"{label}: {value}")
    return "\n\n".join(parts)


def add_comment(doc: Document, paragraph, item: dict[str, Any]) -> None:
    text = comment_text(item)
    if not text or not paragraph.runs:
        return
    doc.add_comment(paragraph.runs, text, author="OpenClaw")


def replace_paragraph_for_field(doc: Document, target_text: str, value: str, item: dict[str, Any]) -> bool:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == target_text.strip():
            set_paragraph_text(paragraph, append_visible_source_block(value, item))
            add_comment(doc, paragraph, item)
            return True
    return False


def replace_prefix_field(doc: Document, target_text: str, value: str, item: dict[str, Any]) -> bool:
    prefix = target_text.rstrip()
    if "：" in prefix:
        prefix = prefix.split("：", 1)[0] + "："
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            filled = f"{prefix}{value}"
            set_paragraph_text(paragraph, append_visible_source_block(filled, item))
            add_comment(doc, paragraph, item)
            return True
    return False


def apply_project_name(doc: Document, company_short_name: str) -> None:
    if not company_short_name:
        return
    for paragraph in doc.paragraphs:
        if "XXXX" in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace("XXXX", company_short_name))


def append_tracking_sections(doc: Document, field_map: dict[str, dict[str, str]], fields: dict[str, dict[str, Any]], sources: list[dict[str, Any]]) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading = doc.add_paragraph()
    heading_run = heading.add_run("字段来源标注")
    heading_run.bold = True
    heading_run.font.size = Pt(14)
    table = doc.add_table(rows=1, cols=6)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    headers = ["字段ID", "字段", "填写内容", "信息来源", "证据等级", "待核实事项"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for fid, item in fields.items():
        meta = field_map.get(fid, {})
        row = table.add_row().cells
        row[0].text = fid
        row[1].text = meta.get("field", "")
        row[2].text = normalize_text(item.get("value"))
        row[3].text = normalize_text(item.get("sources"))
        row[4].text = normalize_text(item.get("evidence_level"))
        row[5].text = normalize_text(item.get("to_verify"))

    heading = doc.add_paragraph()
    heading_run = heading.add_run("来源清单")
    heading_run.bold = True
    heading_run.font.size = Pt(14)
    source_table = doc.add_table(rows=1, cols=5)
    try:
        source_table.style = "Table Grid"
    except KeyError:
        pass
    source_headers = ["编号", "类型", "完整引用", "URL/文件定位", "备注"]
    for i, header in enumerate(source_headers):
        source_table.rows[0].cells[i].text = header
    for source in sources:
        row = source_table.add_row().cells
        row[0].text = normalize_text(source.get("id"))
        row[1].text = normalize_text(source.get("type"))
        row[2].text = normalize_text(source.get("citation"))
        row[3].text = normalize_text(source.get("url_or_locator"))
        row[4].text = normalize_text(source.get("notes"))


def style_document(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10.5)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build source-backed project management interview minutes.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Template .docx path")
    parser.add_argument("--field-map", default=str(DEFAULT_FIELD_MAP), help="Field map CSV path")
    parser.add_argument("--answers", help="Optional JSON answers file")
    parser.add_argument("--output", required=True, help="Output .docx path")
    args = parser.parse_args()

    data: dict[str, Any] = {}
    if args.answers:
        data = json.loads(Path(args.answers).read_text(encoding="utf-8"))

    field_map = load_field_map(Path(args.field_map))
    fields = normalize_fields(data.get("fields"))
    doc = Document(args.template)

    for fid, item in fields.items():
        meta = field_map.get(fid)
        if not meta:
            continue
        target = meta.get("target_text", "")
        value = item.get("value", "")
        if not target:
            continue
        if "：" in target and not target.startswith("（"):
            replace_prefix_field(doc, target, normalize_text(value), item)
        else:
            replace_paragraph_for_field(doc, target, normalize_text(value), item)

    apply_project_name(doc, data.get("company_short_name") or data.get("project_name") or "")
    append_tracking_sections(doc, field_map, fields, data.get("sources", []))
    style_document(doc)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
