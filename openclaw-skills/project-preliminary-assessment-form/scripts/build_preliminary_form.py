#!/usr/bin/env python3
"""Build a source-backed 项目前置研判信息表 workbook from the template."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "附件2：项目前置研判信息表.xlsx"
DEFAULT_FIELD_MAP = SKILL_DIR / "references" / "field-map.csv"


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(x) for x in value if x)
    return str(value)


def load_field_map(path: Path) -> dict[str, dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return {row["field_id"]: row for row in csv.DictReader(f)}


def normalize_fields(raw: Any) -> dict[str, dict[str, Any]]:
    if raw is None:
        return {}
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, list):
        result: dict[str, dict[str, Any]] = {}
        for item in raw:
            fid = item.get("field_id")
            if fid:
                result[fid] = item
        return result
    raise TypeError("fields must be a dict or list")


def add_cell_comment(cell, item: dict[str, Any]) -> None:
    parts = []
    sources = normalize_text(item.get("sources"))
    evidence = normalize_text(item.get("evidence_level"))
    quote = normalize_text(item.get("quote_or_link"))
    verify = normalize_text(item.get("to_verify"))
    if sources:
        parts.append(f"信息来源:\n{sources}")
    if evidence:
        parts.append(f"证据等级: {evidence}")
    if quote:
        parts.append(f"依据摘录/链接:\n{quote}")
    if verify:
        parts.append(f"待核实事项:\n{verify}")
    if parts:
        cell.comment = Comment("\n\n".join(parts), "OpenClaw")


def write_source_tracking_sheet(wb, field_map: dict[str, dict[str, str]], fields: dict[str, dict[str, Any]]) -> None:
    name = "字段来源标注"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    headers = ["字段ID", "工作表", "单元格", "章节", "字段", "填写内容", "信息来源", "证据等级", "依据摘录/链接", "待核实事项"]
    ws.append(headers)
    for fid, item in fields.items():
        meta = field_map.get(fid, {})
        ws.append([
            fid,
            meta.get("sheet", item.get("sheet", "")),
            meta.get("target_cell", item.get("target_cell", "")),
            meta.get("section", ""),
            meta.get("field", item.get("field", "")),
            normalize_text(item.get("value")),
            normalize_text(item.get("sources")),
            normalize_text(item.get("evidence_level")),
            normalize_text(item.get("quote_or_link")),
            normalize_text(item.get("to_verify")),
        ])
    widths = [12, 24, 12, 22, 44, 56, 56, 18, 56, 44]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")


def write_sources_sheet(wb, sources: list[dict[str, Any]]) -> None:
    name = "来源清单"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(["编号", "类型", "完整引用", "URL/文件定位", "备注"])
    for source in sources:
        ws.append([
            source.get("id", ""),
            source.get("type", ""),
            source.get("citation", ""),
            source.get("url_or_locator", ""),
            source.get("notes", ""),
        ])
    widths = {"A": 12, "B": 18, "C": 72, "D": 64, "E": 40}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a source-backed preliminary assessment form.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Template .xlsx path")
    parser.add_argument("--field-map", default=str(DEFAULT_FIELD_MAP), help="Field map CSV path")
    parser.add_argument("--answers", help="Optional JSON answers file")
    parser.add_argument("--output", required=True, help="Output .xlsx path")
    args = parser.parse_args()

    data: dict[str, Any] = {}
    if args.answers:
        data = json.loads(Path(args.answers).read_text(encoding="utf-8"))

    field_map = load_field_map(Path(args.field_map))
    fields = normalize_fields(data.get("fields"))
    wb = load_workbook(args.template)

    for fid, item in fields.items():
        meta = field_map.get(fid)
        sheet_name = item.get("sheet") or (meta or {}).get("sheet")
        target_cell = item.get("target_cell") or (meta or {}).get("target_cell")
        if not sheet_name or not target_cell:
            continue
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        cell = ws[target_cell]
        if "value" in item:
            cell.value = item["value"]
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        add_cell_comment(cell, item)

    write_source_tracking_sheet(wb, field_map, fields)
    write_sources_sheet(wb, data.get("sources", []))

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(output)


if __name__ == "__main__":
    main()
